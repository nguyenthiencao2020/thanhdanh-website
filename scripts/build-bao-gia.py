# Builds tai-lieu/Bao-gia-Dat-Danh.xlsx (quotation template) from src/data/catalog.json.
# Run: python3 scripts/build-bao-gia.py  (then recalculate in Excel/LibreOffice on open).
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

from pathlib import Path
ROOT = str(Path(__file__).resolve().parents[1])
cat = json.load(open(f"{ROOT}/src/data/catalog.json"))
S = cat["settings"]

FONT = "Arial"
RED, BLUE, NAVY, GREY, LIGHT = "B92A26", "2F66A3", "142A44", "64748B", "EFF5FB"
INPUT_FILL = PatternFill("solid", fgColor="FFF8D6")   # cells to fill in
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="CBD5E1")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
VND = '#,##0;-#,##0;"-"'

def f(size=10, bold=False, color="1E293B", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

wb = Workbook()

# ------------------------------------------------------------------ Danh mục (reference lists)
dm = wb.active
dm.title = "Danh mục"
dm["A1"] = "DANH MỤC SẢN PHẨM (giống website)"; dm["A1"].font = f(13, True, NAVY)
dm["A2"] = "Nguồn: cây danh mục trên website (Admin → Danh mục). Sửa ở đây chỉ ảnh hưởng danh sách chọn trong file này."
dm["A2"].font = f(9, italic=True, color=GREY)
hdr = ["Nhóm sản phẩm", "Danh mục con"]
for i, h in enumerate(hdr, 1):
    c = dm.cell(row=4, column=i, value=h); c.font = f(10, True, "FFFFFF"); c.fill = HEAD_FILL; c.border = BOX
groups = [c for c in cat["categories"] if not c.get("parent")]
r = 5
for g in groups:
    kids = [c for c in cat["categories"] if c.get("parent") == g["slug"]]
    for k in kids:
        dm.cell(row=r, column=1, value=g["name"]).font = f()
        dm.cell(row=r, column=2, value=k["name"]).font = f()
        for col in (1, 2): dm.cell(row=r, column=col).border = BOX
        r += 1
sub_last = r - 1
# group list + brand list for dropdowns
dm["D4"] = "Nhóm (danh sách chọn)"; dm["F4"] = "Thương hiệu"
for col in ("D4", "F4"):
    dm[col].font = f(10, True, "FFFFFF"); dm[col].fill = HEAD_FILL; dm[col].border = BOX
for i, g in enumerate(groups):
    c = dm.cell(row=5 + i, column=4, value=g["name"]); c.font = f(); c.border = BOX
brands = [b["name"] for b in cat["brands"]] + ["Khác"]
for i, b in enumerate(brands):
    c = dm.cell(row=5 + i, column=6, value=b); c.font = f(); c.border = BOX
dm.column_dimensions["A"].width = 34; dm.column_dimensions["B"].width = 38
dm.column_dimensions["C"].width = 3; dm.column_dimensions["D"].width = 34
dm.column_dimensions["E"].width = 3; dm.column_dimensions["F"].width = 22
dm.freeze_panes = "A5"
wb.defined_names["DS_Nhom"] = DefinedName("DS_Nhom", attr_text=f"'Danh mục'!$D$5:$D${4 + len(groups)}")
wb.defined_names["DS_DanhMucCon"] = DefinedName("DS_DanhMucCon", attr_text=f"'Danh mục'!$B$5:$B${sub_last}")
wb.defined_names["DS_ThuongHieu"] = DefinedName("DS_ThuongHieu", attr_text=f"'Danh mục'!$F$5:$F${4 + len(brands)}")

# ------------------------------------------------------------------ Bảng giá (price list)
bg = wb.create_sheet("Bảng giá", 0)
bg["A1"] = "BẢNG GIÁ HÀNG HÓA"; bg["A1"].font = f(14, True, RED)
bg["A2"] = "Nhập mỗi mã hàng một dòng. Sheet “Báo giá” tự lấy Tên / Thương hiệu / ĐVT / Đơn giá theo Mã hàng."
bg["A2"].font = f(9, italic=True, color=GREY)
cols = [("Mã hàng", 22), ("Tên hàng / Quy cách", 46), ("Nhóm sản phẩm", 28), ("Danh mục con", 30),
        ("Thương hiệu", 18), ("ĐVT", 9), ("Đơn giá (VNĐ, chưa VAT)", 20), ("Ghi chú", 30)]
for i, (h, w) in enumerate(cols, 1):
    c = bg.cell(row=4, column=i, value=h)
    c.font = f(10, True, "FFFFFF"); c.fill = HEAD_FILL; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bg.column_dimensions[c.column_letter].width = w
bg.row_dimensions[4].height = 30
example = ["6ES7214-1AG40-0XB0", "VÍ DỤ – PLC Siemens S7-1200 CPU 1214C DC/DC/DC (xóa dòng này khi dùng)",
           "Thiết bị tự động hóa (PLC)", "PLC", "Siemens", "Cái", 1000000, "Đơn giá minh họa, không phải giá thật"]
PRICE_ROWS = 500
for r in range(5, 5 + PRICE_ROWS):
    for col in range(1, 9):
        c = bg.cell(row=r, column=col)
        c.border = BOX; c.font = f(); c.fill = INPUT_FILL
        if col == 7: c.number_format = VND
        if col == 2: c.alignment = Alignment(wrap_text=True, vertical="top")
for col, v in enumerate(example, 1):
    c = bg.cell(row=5, column=col, value=v); c.font = f(italic=True, color=GREY)
bg.cell(row=5, column=7).comment = Comment("Đơn giá minh họa để xem định dạng — thay bằng giá thật.", "Đạt Danh")
last_price = 4 + PRICE_ROWS
for rng, name in ((f"C5:C{last_price}", "DS_Nhom"), (f"D5:D{last_price}", "DS_DanhMucCon"), (f"E5:E{last_price}", "DS_ThuongHieu")):
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True, showErrorMessage=False)
    dv.add(rng); bg.add_data_validation(dv)
dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                    showErrorMessage=True, error="Đơn giá phải là số ≥ 0", errorTitle="Đơn giá")
dv.add(f"G5:G{last_price}"); bg.add_data_validation(dv)
bg.freeze_panes = "B5"
bg.auto_filter.ref = f"A4:H{last_price}"
wb.defined_names["BG_Ma"] = DefinedName("BG_Ma", attr_text=f"'Bảng giá'!$A$5:$A${last_price}")
wb.defined_names["BG_Ten"] = DefinedName("BG_Ten", attr_text=f"'Bảng giá'!$B$5:$B${last_price}")
wb.defined_names["BG_Hang"] = DefinedName("BG_Hang", attr_text=f"'Bảng giá'!$E$5:$E${last_price}")
wb.defined_names["BG_DVT"] = DefinedName("BG_DVT", attr_text=f"'Bảng giá'!$F$5:$F${last_price}")
wb.defined_names["BG_Gia"] = DefinedName("BG_Gia", attr_text=f"'Bảng giá'!$G$5:$G${last_price}")

# ------------------------------------------------------------------ Báo giá (quotation)
bq = wb.create_sheet("Báo giá", 0)
widths = {"A": 5, "B": 18, "C": 36, "D": 13, "E": 7, "F": 8, "G": 14, "H": 8, "I": 16, "J": 16}
for k, v in widths.items(): bq.column_dimensions[k].width = v

try:
    img = XLImage(f"{ROOT}/public/emblem.png"); img.height = 74; img.width = 57
    bq.add_image(img, "B1")
except Exception as e:
    print("logo skipped:", e)
bq.merge_cells("C1:J1"); bq["C1"] = S["company_name"].upper(); bq["C1"].font = f(13, True, RED)
bq.merge_cells("C2:J2"); bq["C2"] = (S.get("english_name") or "").upper(); bq["C2"].font = f(8, color=GREY)
bq.merge_cells("C3:J3"); bq["C3"] = f"Địa chỉ: {S['address']}"; bq["C3"].font = f(9)
bq.merge_cells("C4:J4"); bq["C4"] = f"MST: {S.get('tax_code') or ''}   ·   Hotline: {S['phone']} – {S['phone2']}   ·   Email: …………………"
bq["C4"].font = f(9)
for c in ("C1", "C2", "C3", "C4"): bq[c].alignment = Alignment(horizontal="left", vertical="center")
bq.row_dimensions[1].height = 22
for col in range(1, 11):
    bq.cell(row=5, column=col).border = Border(bottom=Side(style="medium", color=RED))

bq.merge_cells("A7:J7"); bq["A7"] = "BẢNG BÁO GIÁ"; bq["A7"].font = f(18, True, NAVY)
bq["A7"].alignment = Alignment(horizontal="center")
bq.merge_cells("A8:J8"); bq["A8"] = '="Số: "&$I$10&"   ·   Ngày: "&TEXT($I$11,"dd/mm/yyyy")'
bq["A8"].font = f(9, italic=True, color=GREY); bq["A8"].alignment = Alignment(horizontal="center")

info = [("Kính gửi:", "B10", "C10:F10"), ("Người liên hệ:", "B11", "C11:F11"), ("Điện thoại:", "B12", "C12:F12"),
        ("Email:", "B13", "C13:F13"), ("Địa chỉ:", "B14", "C14:F14")]
for label, lc, rng in info:
    bq[lc] = label; bq[lc].font = f(10, True)
    bq.merge_cells(rng); first = rng.split(":")[0]
    bq[first].fill = INPUT_FILL; bq[first].font = f(); bq[first].border = BOX
for label, lc, vc in (("Số báo giá:", "H10", "I10"), ("Ngày:", "H11", "I11"), ("Hiệu lực đến:", "H12", "I12")):
    bq[lc] = label; bq[lc].font = f(10, True); bq[lc].alignment = Alignment(horizontal="right")
    bq.merge_cells(f"{vc}:J{vc[1:]}")
    bq[vc].fill = INPUT_FILL; bq[vc].font = f(); bq[vc].border = BOX; bq[vc].alignment = Alignment(horizontal="left")
bq["I10"] = "BG-2026-001"
bq["I11"] = "=TODAY()"; bq["I11"].number_format = "dd/mm/yyyy"
bq["I12"] = "=I11+15"; bq["I12"].number_format = "dd/mm/yyyy"
bq["I12"].comment = Comment("Mặc định 15 ngày kể từ ngày báo giá — sửa nếu cần.", "Đạt Danh")
bq.merge_cells("A16:J16")
bq["A16"] = f"{S['short_name']} xin trân trọng gửi đến Quý khách bảng báo giá như sau:"
bq["A16"].font = f(10, italic=True)

HR = 18
heads = ["STT", "Mã hàng", "Tên hàng / Quy cách", "Thương hiệu", "ĐVT", "SL", "Đơn giá (VNĐ)", "CK (%)", "Thành tiền (VNĐ)", "Ghi chú"]
for i, h in enumerate(heads, 1):
    c = bq.cell(row=HR, column=i, value=h)
    c.font = f(10, True, "FFFFFF"); c.fill = PatternFill("solid", fgColor=RED); c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
bq.row_dimensions[HR].height = 30
N = 25
first, last = HR + 1, HR + N
for i in range(N):
    r = first + i
    bq.cell(row=r, column=1, value=f'=IF(B{r}="","",ROW()-{HR})')
    bq.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(INDEX(BG_Ten,MATCH(B{r},BG_Ma,0)),""))')
    bq.cell(row=r, column=4, value=f'=IF(B{r}="","",IFERROR(INDEX(BG_Hang,MATCH(B{r},BG_Ma,0)),""))')
    bq.cell(row=r, column=5, value=f'=IF(B{r}="","",IFERROR(INDEX(BG_DVT,MATCH(B{r},BG_Ma,0)),""))')
    bq.cell(row=r, column=7, value=f'=IF(B{r}="","",IFERROR(INDEX(BG_Gia,MATCH(B{r},BG_Ma,0)),""))')
    bq.cell(row=r, column=9, value=f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r}*(1-N(H{r})))')
    for col in range(1, 11):
        c = bq.cell(row=r, column=col); c.border = BOX; c.font = f()
        c.alignment = Alignment(vertical="center", wrap_text=(col == 3), horizontal="center" if col in (1, 5, 6, 8) else None)
    for col in (2, 6, 8, 10): bq.cell(row=r, column=col).fill = INPUT_FILL
    bq.cell(row=r, column=7).number_format = VND
    bq.cell(row=r, column=9).number_format = VND
    bq.cell(row=r, column=8).number_format = "0%"
    bq.cell(row=r, column=6).number_format = "#,##0"
dvm = DataValidation(type="list", formula1="=BG_Ma", allow_blank=True, showErrorMessage=False)
dvm.add(f"B{first}:B{last}"); bq.add_data_validation(dvm)
dvq = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                     error="Số lượng phải lớn hơn 0", errorTitle="Số lượng")
dvq.add(f"F{first}:F{last}"); bq.add_data_validation(dvq)
dvc = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True, showErrorMessage=True,
                     error="Nhập % chiết khấu từ 0% đến 100% (VD: 5%)", errorTitle="Chiết khấu")
dvc.add(f"H{first}:H{last}"); bq.add_data_validation(dvc)
bq.cell(row=first, column=2).comment = Comment(
    "Chọn hoặc gõ Mã hàng có trong sheet “Bảng giá” — Tên, Thương hiệu, ĐVT, Đơn giá tự điền.\n"
    "Hàng chưa có trong bảng giá: gõ đè trực tiếp vào các ô Tên / Đơn giá.", "Đạt Danh")

t = last + 1
def total_row(r, label, formula, fmt=VND, bold=False, fill=None, input_cell=None):
    bq.merge_cells(f"A{r}:G{r}")
    bq[f"A{r}"] = label; bq[f"A{r}"].font = f(10, True); bq[f"A{r}"].alignment = Alignment(horizontal="right")
    if input_cell:
        bq[f"H{r}"] = input_cell; bq[f"H{r}"].number_format = "0%"; bq[f"H{r}"].fill = INPUT_FILL
        bq[f"H{r}"].alignment = Alignment(horizontal="center"); bq[f"H{r}"].font = f()
    bq[f"I{r}"] = formula; bq[f"I{r}"].number_format = fmt; bq[f"I{r}"].font = f(11 if bold else 10, True, RED if bold else "1E293B")
    for col in range(1, 11):
        c = bq.cell(row=r, column=col); c.border = BOX
        if fill: c.fill = fill
total_row(t, "Cộng tiền hàng (chưa VAT)", f"=SUM(I{first}:I{last})")
total_row(t + 1, "Thuế GTGT (VAT)", f"=ROUND(I{t}*H{t+1},0)", input_cell=0.1)
bq[f"H{t+1}"].comment = Comment("Thuế suất VAT — mặc định 10%. Kiểm tra mức thuế suất áp dụng cho từng mặt hàng/thời điểm (có thể là 8%).", "Đạt Danh")
total_row(t + 2, "TỔNG CỘNG THANH TOÁN", f"=I{t}+I{t+1}", bold=True, fill=SUB_FILL)
bq.merge_cells(f"A{t+3}:J{t+3}")
bq[f"A{t+3}"] = "Bằng chữ: …………………………………………………………………………………… đồng."
bq[f"A{t+3}"].font = f(10, italic=True); bq[f"A{t+3}"].fill = INPUT_FILL

terms_row = t + 5
bq[f"A{terms_row}"] = "ĐIỀU KHOẢN"; bq[f"A{terms_row}"].font = f(11, True, NAVY)
terms = [
    "1. Giá trên chưa bao gồm thuế GTGT (VAT), đã/không bao gồm chi phí vận chuyển: …………",
    "2. Thời gian giao hàng: …………  ngày làm việc kể từ khi xác nhận đơn hàng.",
    "3. Địa điểm giao hàng: …………",
    "4. Thanh toán: …………  (chuyển khoản / tiền mặt).",
    "5. Bảo hành: theo tiêu chuẩn của nhà sản xuất đối với từng sản phẩm.",
    "6. Báo giá có hiệu lực đến ngày ghi ở trên.",
]
for i, line in enumerate(terms):
    r = terms_row + 1 + i
    bq.merge_cells(f"A{r}:J{r}"); bq[f"A{r}"] = line; bq[f"A{r}"].font = f(9.5)
    if "…" in line: bq[f"A{r}"].fill = INPUT_FILL

sig = terms_row + len(terms) + 2
bq.merge_cells(f"A{sig}:D{sig}"); bq[f"A{sig}"] = "XÁC NHẬN CỦA KHÁCH HÀNG"
bq.merge_cells(f"G{sig}:J{sig}"); bq[f"G{sig}"] = f"ĐẠI DIỆN {S['short_name'].upper()}"
bq.merge_cells(f"A{sig+1}:D{sig+1}"); bq[f"A{sig+1}"] = "(Ký, ghi rõ họ tên)"
bq.merge_cells(f"G{sig+1}:J{sig+1}"); bq[f"G{sig+1}"] = "(Ký, ghi rõ họ tên)"
for c in (f"A{sig}", f"G{sig}"): bq[c].font = f(10, True); bq[c].alignment = Alignment(horizontal="center")
for c in (f"A{sig+1}", f"G{sig+1}"): bq[c].font = f(9, italic=True, color=GREY); bq[c].alignment = Alignment(horizontal="center")

bq.print_area = f"A1:J{sig + 6}"
bq.page_setup.paperSize = bq.PAPERSIZE_A4
bq.page_setup.orientation = "portrait"
bq.page_setup.fitToWidth = 1; bq.page_setup.fitToHeight = 0
bq.sheet_properties.pageSetUpPr.fitToPage = True
bq.page_margins.left = bq.page_margins.right = 0.4
bq.print_title_rows = f"{HR}:{HR}"
bq.oddFooter.center.text = "Trang &P / &N"
bq.sheet_view.showGridLines = False

# example line on the quotation (uses the example price-list row)
bq[f"B{first}"] = "6ES7214-1AG40-0XB0"; bq[f"F{first}"] = 2; bq[f"H{first}"] = 0.05; bq[f"J{first}"] = "Dòng ví dụ – xóa khi dùng"
bq[f"J{first}"].font = f(9, italic=True, color=GREY)

# ------------------------------------------------------------------ Hướng dẫn
hd = wb.create_sheet("Hướng dẫn")
lines = [
    ("HƯỚNG DẪN SỬ DỤNG FILE BÁO GIÁ", f(14, True, RED)),
    ("", None),
    ("Ô nền vàng nhạt = ô cần nhập. Các ô khác là công thức, tự tính.", f(10, True)),
    ("", None),
    ("1. Sheet “Bảng giá”: nhập danh sách hàng — Mã hàng, Tên/quy cách, Nhóm, Danh mục con, Thương hiệu, ĐVT, Đơn giá (chưa VAT).", f()),
    ("   • Nhóm / Danh mục con / Thương hiệu có danh sách chọn, giống danh mục trên website.", f()),
    ("   • Dòng đầu là VÍ DỤ (giá minh họa) — xóa hoặc ghi đè khi dùng.", f()),
    ("2. Sheet “Báo giá”: điền thông tin khách hàng, Số báo giá; Ngày tự lấy hôm nay, Hiệu lực mặc định +15 ngày.", f()),
    ("   • Cột Mã hàng: chọn hoặc gõ mã có trong Bảng giá → Tên, Thương hiệu, ĐVT, Đơn giá tự điền.", f()),
    ("   • Nhập Số lượng và Chiết khấu (%) nếu có → Thành tiền = SL × Đơn giá × (1 − CK).", f()),
    ("   • Hàng chưa có trong Bảng giá: gõ đè trực tiếp vào ô Tên / Thương hiệu / ĐVT / Đơn giá của dòng đó.", f()),
    ("   • Thuế VAT mặc định 10% (ô vàng cạnh dòng VAT) — kiểm tra thuế suất áp dụng, có thể là 8%.", f()),
    ("   • Điền “Bằng chữ” và các điều khoản có dấu …… trước khi gửi.", f()),
    ("3. In / xuất PDF: File → Print (đã cài sẵn khổ A4 dọc, vừa 1 trang ngang).", f()),
    ("4. Mỗi báo giá: File → Save As với tên mới, VD: BG-2026-001_TenKhachHang.xlsx.", f()),
    ("", None),
    ("Bảng báo giá có tối đa 25 dòng hàng; Bảng giá có sẵn 500 dòng. Thông tin công ty lấy theo giấy chứng nhận ĐKDN (MST 0319753015).", f(9, italic=True, color=GREY)),
]
for i, (txt, font) in enumerate(lines, 1):
    c = hd.cell(row=i, column=1, value=txt)
    if font: c.font = font
hd["A3"].fill = INPUT_FILL
hd.column_dimensions["A"].width = 120

wb.active = 0
out = f"{ROOT}/tai-lieu/Bao-gia-Dat-Danh.xlsx"
wb.save(out)
print("saved", out)
